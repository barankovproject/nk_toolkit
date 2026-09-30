from __future__ import annotations

import clr

clr.AddReference("AecPropDataMgd")
clr.AddReference("acdbmgd")

import json
import os
import traceback
from collections import defaultdict
from contextlib import contextmanager
from typing import Any

import Autodesk.Aec.PropertyData.DatabaseServices as _PD
from Autodesk.AutoCAD.ApplicationServices import Application
from Autodesk.AutoCAD.DatabaseServices import OpenMode, Solid3d

from civil.utils import normalize
from paths import GSI_BASKETS_FILE, REPORT_CONFIG_FILE, REPORTS_DIR, TYPES_FILE

from .logger import Logger


def _load_report_config() -> dict[str, Any]:
    with open(REPORT_CONFIG_FILE, encoding="utf-8") as f:
        return json.load(f)["rows"]


def _load_canal_types() -> dict[str, Any]:
    try:
        with open(TYPES_FILE, encoding="utf-8-sig") as f:
            return json.load(f).get("types", {})
    except Exception:
        return {}


def _load_gsi_baskets() -> dict[str, str]:
    try:
        with open(GSI_BASKETS_FILE, encoding="utf-8-sig") as f:
            return json.load(f).get("baskets", {})
    except Exception:
        return {}


@contextmanager
def _read_tx(db: Any):
    tx = db.TransactionManager.StartTransaction()
    try:
        yield tx
    finally:
        tx.Abort()
        tx.Dispose()


def _psd_id(dpsd: Any, name: str) -> Any:
    return dpsd.GetAt(name) if name in list(dpsd.NamesInUse) else None


def _ps_read(tx: Any, ent: Any, psd_id: Any) -> Any:
    try:
        ps_id = _PD.PropertyDataServices.GetPropertySet(ent, psd_id)
        return tx.GetObject(ps_id, OpenMode.ForRead)
    except Exception:
        return None


def _get(ps: Any, field: str, default: Any = None) -> Any:
    try:
        return ps.GetAt(ps.PropertyNameToId(field))
    except Exception:
        return default


class ReportBuilder:
    """Reads property sets from the active drawing and writes a quantities Excel report."""

    def __init__(self, query: str, log_keep: int = 10) -> None:
        self.query = query.strip()
        self.log = Logger(keep=log_keep)

    def run(self) -> str:
        log = self.log
        log("=== CANAL REPORT STARTED ===")
        log(f"query='{self.query}'")

        try:
            cfg = _load_report_config()
            canal_types = _load_canal_types()
            gsi_baskets = _load_gsi_baskets()

            doc = Application.DocumentManager.MdiActiveDocument
            db = doc.Database

            with _read_tx(db) as tx:
                dpsd = _PD.DictionaryPropertySetDefinitions(db)

                psd_ids: dict[str, Any] = {}
                for row in cfg.values():
                    psd_name = row["source_psd"]
                    if psd_name not in psd_ids:
                        psd_ids[psd_name] = _psd_id(dpsd, psd_name)

                gabion_id = psd_ids.get("Arhyz_Gabion")
                rib_id = psd_ids.get("Arhyz_Rib")

                if gabion_id is None and rib_id is None:
                    raise Exception(
                        "No Arhyz property sets found in drawing — run canal_model first"
                    )

                canal_name = self._find_canal_name(tx, db, gabion_id, rib_id)
                log(f"matched canal: '{canal_name}'")

                # accumulators keyed by row id
                scalars: dict[str, float] = defaultdict(float)
                basket_counts: dict[str, int] = defaultdict(int)

                ms = tx.GetObject(db.CurrentSpaceId, OpenMode.ForRead)
                for eid in ms:
                    try:
                        ent = tx.GetObject(eid, OpenMode.ForRead)
                        if not isinstance(ent, Solid3d):
                            continue

                        gabion_matched = False
                        type_key = ""

                        for row_id, row in cfg.items():
                            psd_name = row["source_psd"]
                            # basket row: data comes from canal_types config, not PS
                            if psd_name == "Arhyz_Gabion_Baskets":
                                continue
                            pid = psd_ids.get(psd_name)
                            if pid is None:
                                continue
                            ps = _ps_read(tx, ent, pid)
                            if ps is None:
                                continue
                            if str(_get(ps, "Canal", "") or "") != canal_name:
                                continue

                            if psd_name == "Arhyz_Gabion":
                                gabion_matched = True
                                if not type_key:
                                    type_key = str(_get(ps, "Type", "") or "")

                            val = _get(ps, row["source_field"], 0.0)
                            scalars[row_id] += float(val or 0.0)

                        # accumulate basket count from config using Type field
                        if gabion_matched and type_key in canal_types:
                            tp = canal_types[type_key]
                            dim_key = str(tp.get("basket_dim", ""))
                            gsi_name = gsi_baskets.get(dim_key, dim_key)
                            b_cnt = int(tp.get("basket_count", 0))
                            if gsi_name:
                                basket_counts[gsi_name] += b_cnt

                    except Exception:
                        continue

            log(f"scalars: { {k: round(v, 4) for k, v in scalars.items()} }")
            log(f"baskets: {dict(basket_counts)}")

            out_path = self._write_excel(canal_name, cfg, scalars, basket_counts)
            log(f"report saved: {out_path}")

        except Exception as ex:
            log(f"ERROR: {ex}", "CRITICAL")
            log(traceback.format_exc(), "CRITICAL")
            return f"ERROR: {ex}"

        log("=== DONE ===")
        return out_path

    def _find_canal_name(self, tx: Any, db: Any, gabion_id: Any, rib_id: Any) -> str:
        """Return the exact Canal PS value that best matches self.query."""
        q = normalize(self.query)
        seen: set[str] = set()
        ms = tx.GetObject(db.CurrentSpaceId, OpenMode.ForRead)
        for eid in ms:
            try:
                ent = tx.GetObject(eid, OpenMode.ForRead)
                if not isinstance(ent, Solid3d):
                    continue
                for psd_id in (gabion_id, rib_id):
                    if psd_id is None:
                        continue
                    ps = _ps_read(tx, ent, psd_id)
                    if ps is not None:
                        name = str(_get(ps, "Canal", "") or "")
                        if name:
                            seen.add(name)
            except Exception:
                continue

        if not seen:
            raise Exception("No tagged solids found — run canal_model first")

        for name in seen:
            if normalize(name) == q:
                return name
        sw = [n for n in seen if normalize(n).startswith(q)]
        if sw:
            return min(sw, key=len)
        co = [n for n in seen if q in normalize(n)]
        if co:
            return min(co, key=len)
        raise Exception(
            f"No canal matching '{self.query}' in drawing. Available: {sorted(seen)}"
        )

    def _write_excel(
        self,
        canal_name: str,
        cfg: dict[str, Any],
        scalars: dict[str, float],
        basket_counts: dict[str, int],
    ) -> str:
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill

            return self._write_xlsx(
                canal_name,
                cfg,
                scalars,
                basket_counts,
                openpyxl,
                Alignment,
                Font,
                PatternFill,
            )
        except ImportError:
            self.log("openpyxl not found — falling back to CSV", "WARN")
            return self._write_csv(canal_name, cfg, scalars, basket_counts)

    def _build_rows(
        self,
        cfg: dict[str, Any],
        scalars: dict[str, float],
        basket_counts: dict[str, int],
    ) -> list[tuple[str, str, Any]]:
        rows: list[tuple[str, str, Any]] = []
        for row_id, row in cfg.items():
            if row["source_psd"] == "Arhyz_Gabion_Baskets":
                # expand one row per basket dimension
                for dim in sorted(basket_counts):
                    label = f"{row['name_prefix']} {dim}"
                    rows.append((label, row["unit"], basket_counts[dim]))
            else:
                val = scalars.get(row_id, 0.0)
                qty = val * row["coef"]
                fmt_qty = round(qty, 3) if row["unit"] == "м³" else round(qty, 2)
                rows.append((row["name"], row["unit"], fmt_qty))
        return rows

    def _write_xlsx(
        self,
        canal_name: str,
        cfg: dict[str, Any],
        scalars: dict[str, float],
        basket_counts: dict[str, int],
        openpyxl: Any,
        Alignment: Any,
        Font: Any,
        PatternFill: Any,
    ) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ведомость объемов"

        header_fill = PatternFill("solid", fgColor="D9E1F2")
        header_font = Font(bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 70
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 14

        headers = ["№", "Наименование работ", "Ед. изм.", "Количество"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        rows = self._build_rows(cfg, scalars, basket_counts)
        for i, (name, unit, qty) in enumerate(rows, 1):
            r = i + 1
            ws.cell(row=r, column=1, value=i).alignment = center
            ws.cell(row=r, column=2, value=name).alignment = left
            ws.cell(row=r, column=3, value=unit).alignment = center
            ws.cell(row=r, column=4, value=qty).alignment = center

        ws.row_dimensions[1].height = 30

        os.makedirs(REPORTS_DIR, exist_ok=True)
        out_path = os.path.join(REPORTS_DIR, f"{canal_name}.xlsx")
        try:
            wb.save(out_path)
        except PermissionError:
            import datetime

            ts = datetime.datetime.now().strftime("%H%M%S")
            out_path = os.path.join(REPORTS_DIR, f"{canal_name}_{ts}.xlsx")
            wb.save(out_path)
        return out_path

    def _write_csv(
        self,
        canal_name: str,
        cfg: dict[str, Any],
        scalars: dict[str, float],
        basket_counts: dict[str, int],
    ) -> str:
        import csv

        rows = self._build_rows(cfg, scalars, basket_counts)
        out_rows = [["№", "Наименование работ", "Ед. изм.", "Количество"]]
        for i, (name, unit, qty) in enumerate(rows, 1):
            out_rows.append([i, name, unit, qty])

        os.makedirs(REPORTS_DIR, exist_ok=True)
        out_path = os.path.join(REPORTS_DIR, f"{canal_name}.csv")
        try:
            with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f, delimiter=";").writerows(out_rows)
        except PermissionError:
            import datetime

            ts = datetime.datetime.now().strftime("%H%M%S")
            out_path = os.path.join(REPORTS_DIR, f"{canal_name}_{ts}.csv")
            with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f, delimiter=";").writerows(out_rows)
        return out_path
